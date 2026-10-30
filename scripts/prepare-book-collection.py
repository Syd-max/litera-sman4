from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\mirsy\Downloads\Salinan dari DAFTAR KOLEKSI NAIK DISPLAY.xlsx")
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "tmp"
IMPORT_FILE = OUTPUT_DIR / "litera-book-collection-import.tsv"
AUDIT_FILE = OUTPUT_DIR / "litera-book-collection-audit.json"

CATEGORIES = {
    "PENGETAHUAN UMUM": "Pengetahuan Umum",
    "FILSAFAH dan PSIKOLOGI": "Filsafah dan Psikologi",
    "AGAMA": "Agama",
    "ILMU SOSIAL": "Ilmu Sosial",
    "BAHASA": "Bahasa",
    "ILMU MURNI": "Ilmu Murni",
    "ILMU TERAPAN": "Ilmu Terapan",
    "OLAHRAGA dan REKREASI": "Olahraga dan Rekreasi",
    "LITERATUR": "Literatur",
    "GEOGRAFI dan SEJARAH": "Geografi dan Sejarah",
}

HEADERS = [
    "source file", "source sheet", "source row", "no", "nomor induk", "judul", "pengarang",
    "klasifikasi", "penerbit", "tahun terbit", "eksemplar", "sumber", "keterangan", "rak", "kategori",
]


def text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        number_format = str(cell.number_format or "")
        if re.fullmatch(r"0+", number_format):
            return str(int(value)).zfill(len(number_format))
        return str(int(value))
    return str(value).strip()


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    rows: list[dict[str, str]] = []
    orphan_rows: list[dict[str, object]] = []
    category = ""

    for sheet in workbook.worksheets:
        for row_number in range(1, sheet.max_row + 1):
            cells = [sheet.cell(row_number, column) for column in range(1, 12)]
            values = [text(cell) for cell in cells]
            if values[0] in CATEGORIES:
                category = CATEGORIES[values[0]]
                continue
            if row_number < 17 or values[0].upper() == "NO" or not any(values):
                continue
            if not values[2]:
                orphan_rows.append({"sheet": sheet.title, "row": row_number, "values": values})
                continue
            rows.append(dict(zip(HEADERS, [SOURCE.name, sheet.title, str(row_number), *values[:11], category], strict=True)))

    exact_groups: dict[tuple[str, ...], list[int]] = defaultdict(list)
    inventory_groups: dict[str, list[int]] = defaultdict(list)
    for item in rows:
        exact_groups[(
            normalized(item["judul"]), normalized(item["pengarang"]), normalized(item["penerbit"]),
            item["tahun terbit"], normalized(item["klasifikasi"]),
        )].append(int(item["source row"]))
        if item["nomor induk"]:
            inventory_groups[item["nomor induk"]].append(int(item["source row"]))

    OUTPUT_DIR.mkdir(exist_ok=True)
    with IMPORT_FILE.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, dialect="excel-tab")
        writer.writeheader()
        writer.writerows(rows)

    missing = {
        field: sum(not item[field] for item in rows)
        for field in ["nomor induk", "pengarang", "klasifikasi", "penerbit", "tahun terbit", "eksemplar", "sumber", "keterangan", "rak"]
    }
    audit = {
        "source": str(SOURCE),
        "sheets": workbook.sheetnames,
        "book_rows": len(rows),
        "orphan_rows": orphan_rows,
        "categories": dict(Counter(item["kategori"] for item in rows)),
        "missing": missing,
        "exact_duplicate_groups": [group for group in exact_groups.values() if len(group) > 1],
        "repeated_inventory_numbers": [
            {"inventory_number": number, "rows": group}
            for number, group in inventory_groups.items() if len(group) > 1
        ],
    }
    AUDIT_FILE.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")

    assert len(rows) == 3543, f"Expected 3543 book rows, found {len(rows)}"
    assert sum(audit["categories"].values()) == len(rows)
    print(json.dumps({"import_file": str(IMPORT_FILE), "audit_file": str(AUDIT_FILE), "book_rows": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise
